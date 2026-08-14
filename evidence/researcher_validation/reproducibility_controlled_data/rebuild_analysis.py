#!/usr/bin/env python3
"""Rebuild the formal researcher-validation analysis from archived raw sources."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import statistics
import sys
import tempfile
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

import openpyxl


PARTICIPANT_NAMES = {"P01": "P01", "P02": "P02", "P03": "P03"}
FORMAL_PARTICIPANTS = tuple(PARTICIPANT_NAMES)
TOPIC_ROWS_PER_RUN = 4
ITEM_ROWS_PER_RUN = 6


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalise_token(value: Any) -> str:
    text = clean_text(value).casefold()
    text = text.replace("—", "-").replace("–", "-")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def normalise_status(value: Any) -> str:
    token = normalise_token(value)
    if "partial" in token:
        return "partially_answered"
    if "fully" in token or token == "covered":
        return "fully_answered"
    if "not_answered" in token or "not_covered" in token:
        return "not_answered"
    return token


def normalise_action(value: Any) -> str:
    token = normalise_token(value)
    if "exclude" in token:
        return "exclude"
    if "flag" in token and ("edit" in token or "partial" in token or "missing" in token):
        return "edit_and_flag_partial"
    if "correct_source" in token or "edit_source" in token:
        return "edit_source"
    if "edit" in token or "narrow" in token:
        return "edit"
    if "include" in token or token == "included":
        return "include"
    return token


def split_sources(value: Any) -> set[str]:
    return set(re.findall(r"\b[AB]-M\d{2}\b", clean_text(value), flags=re.IGNORECASE))


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_workbook_bytes(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)


def parse_response_workbook(data: bytes, member_name: str) -> dict[str, Any]:
    workbook = load_workbook_bytes(data)
    sheet = workbook["Post-Condition"]
    participant = clean_text(sheet["B4"].value)
    round_number = int(sheet["D4"].value)
    condition = clean_text(sheet["B5"].value)
    case = clean_text(sheet["D5"].value)
    responses: dict[str, Any] = {}
    for row in range(8, 19):
        response_id = clean_text(sheet.cell(row, 1).value)
        if re.fullmatch(r"PC\d+", response_id):
            responses[response_id] = sheet.cell(row, 3).value

    comparison: dict[str, Any] = {}
    if "Final Comparison" in workbook.sheetnames:
        comparison_sheet = workbook["Final Comparison"]
        for row in range(8, 18):
            response_id = clean_text(comparison_sheet.cell(row, 1).value)
            if re.fullmatch(r"D\d+", response_id):
                comparison[response_id] = comparison_sheet.cell(row, 3).value

    if participant not in FORMAL_PARTICIPANTS:
        raise ValueError(f"Unexpected formal participant in {member_name}: {participant!r}")
    if condition not in {"Manual", "Prototype"}:
        raise ValueError(f"Unexpected condition in {member_name}: {condition!r}")
    return {
        "participant": participant,
        "name": PARTICIPANT_NAMES[participant],
        "round": round_number,
        "condition": condition,
        "case": case,
        "protocol": case.split("-")[0],
        "source_member": member_name,
        **responses,
        "comparison": comparison,
    }


def parse_manual_workbook(data: bytes, member_name: str) -> dict[str, Any]:
    workbook = load_workbook_bytes(data)
    topic_sheet = workbook["4 Topic Review"]
    item_sheet = workbook["5 Evidence Review"]
    case_match = re.search(r"NP\d{2}-[AB]", clean_text(topic_sheet["A1"].value))
    if not case_match:
        raise ValueError(f"Cannot identify case in manual workbook {member_name}")
    case = case_match.group(0)

    topics: list[dict[str, Any]] = []
    for row in range(5, 9):
        topic_label = clean_text(topic_sheet.cell(row, 1).value)
        match = re.search(r"\(([^()]+)\)", topic_label)
        if not match:
            continue
        topics.append(
            {
                "case": case,
                "topic_code": match.group(1),
                "actual_status": normalise_status(topic_sheet.cell(row, 3).value),
                "actual_missing_information": clean_text(topic_sheet.cell(row, 4).value),
                "source_file": member_name,
            }
        )

    items: list[dict[str, Any]] = []
    for row in range(5, 11):
        item_id = clean_text(item_sheet.cell(row, 1).value)
        if not re.fullmatch(r"NP\d{2}-[AB]-E\d", item_id):
            continue
        missingness_result = clean_text(item_sheet.cell(row, 9).value)
        missing_information = clean_text(item_sheet.cell(row, 10).value)
        items.append(
            {
                "case": case,
                "item_id": item_id,
                "actual_support": clean_text(item_sheet.cell(row, 5).value),
                "actual_action": normalise_action(item_sheet.cell(row, 6).value),
                "actual_sources": sorted(split_sources(item_sheet.cell(row, 7).value)),
                "actual_missingness": (
                    "partial" in normalise_token(missingness_result)
                    or bool(missing_information and normalise_token(missing_information) not in {"none", "not_applicable"})
                ),
                "actual_missing_information": missing_information,
                "source_file": member_name,
            }
        )
    if len(topics) != TOPIC_ROWS_PER_RUN or len(items) != ITEM_ROWS_PER_RUN:
        raise ValueError(
            f"Manual workbook {member_name} yielded {len(topics)} topic rows and {len(items)} item rows"
        )
    return {"case": case, "topics": topics, "items": items}


class SummaryTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[dict[str, Any]]]] = []
        self._in_table = False
        self._table_depth = 0
        self._current_table: list[list[dict[str, Any]]] = []
        self._current_row: list[dict[str, Any]] | None = None
        self._current_cell: dict[str, Any] | None = None
        self._current_link_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = dict(attrs)
        if tag == "table":
            if not self._in_table:
                self._in_table = True
                self._table_depth = 1
                self._current_table = []
            else:
                self._table_depth += 1
        elif self._in_table and tag == "tr" and self._table_depth == 1:
            self._current_row = []
        elif self._in_table and tag in {"td", "th"} and self._table_depth == 1:
            self._current_cell = {"parts": [], "links": [], "tag": tag}
        elif self._current_cell is not None and tag == "a":
            self._current_link_parts = []
            href = attrs_dict.get("href")
            if href:
                self._current_cell.setdefault("hrefs", []).append(href)

    def handle_data(self, data: str) -> None:
        if self._current_cell is not None:
            self._current_cell["parts"].append(data)
            if self._current_link_parts is not None:
                self._current_link_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._current_cell is not None and self._current_link_parts is not None:
            link_text = clean_text("".join(self._current_link_parts))
            if link_text:
                self._current_cell["links"].append(link_text)
            self._current_link_parts = None
        elif self._in_table and tag in {"td", "th"} and self._current_cell is not None:
            cell = {
                "text": clean_text("".join(self._current_cell["parts"])),
                "links": list(self._current_cell.get("links", [])),
                "hrefs": list(self._current_cell.get("hrefs", [])),
                "tag": self._current_cell["tag"],
            }
            if self._current_row is not None:
                self._current_row.append(cell)
            self._current_cell = None
        elif self._in_table and tag == "tr" and self._current_row is not None and self._table_depth == 1:
            if self._current_row:
                self._current_table.append(self._current_row)
            self._current_row = None
        elif tag == "table" and self._in_table:
            self._table_depth -= 1
            if self._table_depth == 0:
                if self._current_table:
                    self.tables.append(self._current_table)
                self._current_table = []
                self._in_table = False


def source_ids_from_prototype_cell(item_id: str, cell: dict[str, Any]) -> list[str]:
    case_letter = item_id.split("-")[1]
    response_numbers: list[int] = []
    for link in cell.get("links", []):
        match = re.search(r"Participant response\s+(\d+)", link, flags=re.IGNORECASE)
        if match:
            response_numbers.append(int(match.group(1)))
    source_ids = [f"{case_letter}-M{(number - 1) * 2:02d}" for number in response_numbers if number >= 2]
    return sorted(set(source_ids))


def parse_prototype_html(data: bytes, member_name: str) -> dict[str, Any]:
    parser = SummaryTableParser()
    parser.feed(data.decode("utf-8-sig"))
    items: dict[str, dict[str, Any]] = {}
    case = ""
    for table in parser.tables:
        if not table:
            continue
        header = [cell["text"] for cell in table[0]]
        if len(header) >= 3 and header[:3] == ["Topic", "Included extract", "Source / review"]:
            for row in table[1:]:
                if len(row) < 3:
                    continue
                match = re.search(r"NP\d{2}-[AB]-E\d", row[0]["text"])
                if not match:
                    continue
                item_id = match.group(0)
                case = item_id.rsplit("-", 1)[0]
                decision_text = row[2]["text"]
                action = "edit" if decision_text.startswith("Edited") else "include"
                status = "partially_answered" if "Partially covered" in decision_text else "fully_answered"
                items[item_id] = {
                    "case": case,
                    "item_id": item_id,
                    "actual_action": action,
                    "actual_sources": source_ids_from_prototype_cell(item_id, row[2]),
                    "actual_status": status,
                    "actual_missingness": status == "partially_answered",
                    "actual_missing_information": "Protocol-required information flagged" if status == "partially_answered" else "",
                    "source_file": member_name,
                }
        elif len(header) >= 2 and header[:2] == ["Topic", "Exclusion reason"]:
            for row in table[1:]:
                if len(row) < 2:
                    continue
                match = re.search(r"NP\d{2}-[AB]-E\d", row[0]["text"])
                if not match:
                    continue
                item_id = match.group(0)
                case = item_id.rsplit("-", 1)[0]
                items[item_id] = {
                    "case": case,
                    "item_id": item_id,
                    "actual_action": "exclude",
                    "actual_sources": [],
                    "actual_status": "",
                    "actual_missingness": False,
                    "actual_missing_information": "",
                    "source_file": member_name,
                }
    if not case or len(items) != ITEM_ROWS_PER_RUN:
        raise ValueError(f"Prototype HTML {member_name} yielded case={case!r}, item rows={len(items)}")
    return {"case": case, "items": list(items.values())}


def parse_gold(gold_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(gold_path, data_only=True, read_only=True)
    topic_sheet = workbook["Topic Gold Review"]
    item_sheet = workbook["Item Gold Review"]

    topics: list[dict[str, Any]] = []
    for row in range(5, 29):
        case = clean_text(topic_sheet.cell(row, 1).value)
        if not case:
            continue
        decision = clean_text(topic_sheet.cell(row, 7).value)
        row_qa = clean_text(topic_sheet.cell(row, 11).value)
        if decision != "Confirm" or row_qa != "PASS":
            raise ValueError(f"Gold topic row {row} is not a confirmed PASS row")
        topics.append(
            {
                "case": case,
                "topic_code": clean_text(topic_sheet.cell(row, 2).value),
                "topic": clean_text(topic_sheet.cell(row, 3).value),
                "gold_status": normalise_status(topic_sheet.cell(row, 5).value),
                "gold_missing_information": clean_text(topic_sheet.cell(row, 6).value),
            }
        )

    items: list[dict[str, Any]] = []
    for row in range(5, 41):
        case = clean_text(item_sheet.cell(row, 1).value)
        item_id = clean_text(item_sheet.cell(row, 2).value)
        if not case or not item_id:
            continue
        decision = clean_text(item_sheet.cell(row, 14).value)
        row_qa = clean_text(item_sheet.cell(row, 18).value)
        if decision != "Confirm" or row_qa != "PASS":
            raise ValueError(f"Gold item row {row} is not a confirmed PASS row")
        items.append(
            {
                "case": case,
                "item_id": item_id,
                "topic_code": clean_text(item_sheet.cell(row, 3).value),
                "gold_support": clean_text(item_sheet.cell(row, 6).value),
                "gold_sources": sorted(split_sources(item_sheet.cell(row, 7).value)),
                "gold_action": normalise_action(item_sheet.cell(row, 8).value),
                "designed_issue": clean_text(item_sheet.cell(row, 9).value),
                "critical_error_rule": clean_text(item_sheet.cell(row, 12).value),
                "full_score_rule": clean_text(item_sheet.cell(row, 13).value),
            }
        )
    if len(topics) != 24 or len(items) != 36:
        raise ValueError(f"Gold yielded {len(topics)} topic rows and {len(items)} item rows")
    return topics, items


def load_scoring_rules(script_dir: Path) -> dict[str, Any]:
    with (script_dir / "scoring_rules.json").open("r", encoding="utf-8") as handle:
        return json.load(handle)


def item_is_correct(actual: dict[str, Any], gold: dict[str, Any], rules: dict[str, Any]) -> bool:
    expected = gold["gold_action"]
    actual_action = actual["actual_action"]
    rule = rules["item_action_rules"][expected]
    if rule.get("allow_any_non_exclude"):
        action_ok = actual_action != "exclude"
    else:
        action_ok = actual_action in set(rule["allowed_actual_actions"])
    if not action_ok:
        return False
    if rule.get("require_missingness") and not actual.get("actual_missingness", False):
        return False
    if rule.get("require_correct_sources") and actual_action != "exclude":
        expected_sources = set(gold["gold_sources"])
        actual_sources = set(actual.get("actual_sources", []))
        if not expected_sources.issubset(actual_sources):
            return False
    return True


def is_false_approval(actual: dict[str, Any], gold: dict[str, Any]) -> bool:
    expected = gold["gold_action"]
    actual_action = actual["actual_action"]
    if expected == "exclude":
        return actual_action != "exclude"
    if expected == "edit":
        return actual_action == "include"
    if expected == "edit_source":
        expected_sources = set(gold["gold_sources"])
        actual_sources = set(actual.get("actual_sources", []))
        return actual_action != "exclude" and not expected_sources.issubset(actual_sources)
    if expected == "edit_and_flag_partial":
        return actual_action != "exclude" and not actual.get("actual_missingness", False)
    return False


def mean(values: Iterable[float]) -> float:
    values = list(values)
    return sum(values) / len(values)


def build_analysis(package_root: Path) -> dict[str, Any]:
    package_root = package_root.resolve()
    script_dir = Path(__file__).resolve().parent
    rules = load_scoring_rules(script_dir)
    raw_dir = package_root / "01_正式研究_Formal" / "01_原始包_Raw"
    gold_path = package_root / "03_Gold_GoldReviewer01" / "03_最终Gold_Final_Gold_v0.8.0.xlsx"
    raw_zips = sorted(raw_dir.glob("*.zip"))
    if len(raw_zips) != 3:
        raise ValueError(f"Expected three formal raw ZIPs, found {len(raw_zips)}")
    if not gold_path.exists():
        raise FileNotFoundError(gold_path)

    responses: list[dict[str, Any]] = []
    manual_by_case: dict[str, dict[str, Any]] = {}
    prototype_by_case: dict[str, dict[str, Any]] = {}
    input_hashes: list[dict[str, str]] = []

    for raw_zip in raw_zips:
        participant_match = re.search(r"P0[123]", raw_zip.name)
        if not participant_match:
            raise ValueError(f"Cannot identify participant from {raw_zip.name}")
        participant = participant_match.group(0)
        input_hashes.append(
            {
                "input_type": "formal_raw_zip",
                "participant": participant,
                "path": raw_zip.relative_to(package_root).as_posix(),
                "sha256": sha256_file(raw_zip),
            }
        )
        with zipfile.ZipFile(raw_zip) as archive:
            for member in sorted(name for name in archive.namelist() if not name.endswith("/")):
                data = archive.read(member)
                basename = Path(member).name
                input_hashes.append(
                    {
                        "input_type": "archived_member",
                        "participant": participant,
                        "path": f"{raw_zip.relative_to(package_root).as_posix()}::{member}",
                        "sha256": sha256_bytes(data),
                    }
                )
                lower = basename.casefold()
                if lower.endswith(".xlsx") and "response_completed" in lower:
                    responses.append(parse_response_workbook(data, basename))
                elif lower.endswith(".xlsx") and "manual" in lower:
                    parsed = parse_manual_workbook(data, basename)
                    manual_by_case[parsed["case"]] = parsed
                elif lower.endswith(".html") and "evidence_record" in lower:
                    parsed = parse_prototype_html(data, basename)
                    prototype_by_case[parsed["case"]] = parsed

    input_hashes.append(
        {
            "input_type": "gold_authority",
            "participant": "GoldReviewer01",
            "path": gold_path.relative_to(package_root).as_posix(),
            "sha256": sha256_file(gold_path),
        }
    )
    gold_topics, gold_items = parse_gold(gold_path)

    responses.sort(key=lambda row: (row["participant"], row["round"]))
    run_by_case = {row["case"]: row for row in responses}
    if len(responses) != 6 or len(run_by_case) != 6:
        raise ValueError(f"Expected six unique response runs, found {len(responses)} / {len(run_by_case)}")
    if set(manual_by_case) != {row["case"] for row in responses if row["condition"] == "Manual"}:
        raise ValueError("Manual cases do not match the Manual response records")
    if set(prototype_by_case) != {row["case"] for row in responses if row["condition"] == "Prototype"}:
        raise ValueError("Prototype cases do not match the Prototype response records")

    gold_item_by_id = {row["item_id"]: row for row in gold_items}
    gold_topics_by_case: dict[str, list[dict[str, Any]]] = defaultdict(list)
    gold_items_by_case: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in gold_topics:
        gold_topics_by_case[row["case"]].append(row)
    for row in gold_items:
        gold_items_by_case[row["case"]].append(row)

    actual_topics: dict[tuple[str, str], dict[str, Any]] = {}
    actual_items: dict[str, dict[str, Any]] = {}
    for case, manual in manual_by_case.items():
        for row in manual["topics"]:
            actual_topics[(case, row["topic_code"])] = row
        for row in manual["items"]:
            actual_items[row["item_id"]] = row
    for case, prototype in prototype_by_case.items():
        for row in prototype["items"]:
            row["topic_code"] = gold_item_by_id[row["item_id"]]["topic_code"]
            actual_items[row["item_id"]] = row
        for gold_topic in gold_topics_by_case[case]:
            included_statuses = [
                row["actual_status"]
                for row in prototype["items"]
                if row.get("topic_code") == gold_topic["topic_code"] and row["actual_action"] != "exclude"
            ]
            if "partially_answered" in included_statuses:
                actual_status = "partially_answered"
            elif included_statuses:
                actual_status = "fully_answered"
            else:
                actual_status = "not_answered"
            actual_topics[(case, gold_topic["topic_code"])] = {
                "case": case,
                "topic_code": gold_topic["topic_code"],
                "actual_status": actual_status,
                "actual_missing_information": "Protocol-required information flagged" if actual_status == "partially_answered" else "",
                "source_file": prototype["items"][0]["source_file"],
            }

    objective_rows: list[dict[str, Any]] = []
    for gold in gold_topics:
        run = run_by_case[gold["case"]]
        actual = actual_topics[(gold["case"], gold["topic_code"])]
        missing_expected = bool(gold["gold_missing_information"])
        actual_missing_token = normalise_token(actual["actual_missing_information"])
        missing_actual = (
            actual["actual_status"] == "partially_answered"
            or actual_missing_token not in {"", "none", "not_applicable"}
        )
        correct = actual["actual_status"] == gold["gold_status"]
        objective_rows.append(
            {
                "row_id": f"{gold['case']}-TOPIC-{gold['topic_code']}",
                "participant": run["participant"],
                "condition": run["condition"],
                "round": run["round"],
                "case": gold["case"],
                "level": "Topic",
                "target_id": gold["topic_code"],
                "topic_code": gold["topic_code"],
                "gold_decision": gold["gold_status"],
                "actual_decision": actual["actual_status"],
                "correct": correct,
                "issue_bearing": False,
                "issue_recognized": "",
                "missingness_expected": missing_expected,
                "missingness_actual": missing_actual,
                "missingness_correct": (missing_expected == missing_actual) if missing_expected else "",
                "false_approval": False,
                "source_file": actual["source_file"],
            }
        )

    for gold in gold_items:
        run = run_by_case[gold["case"]]
        actual = actual_items[gold["item_id"]]
        correct = item_is_correct(actual, gold, rules)
        issue_bearing = normalise_token(gold["designed_issue"]) not in {"", "none"}
        objective_rows.append(
            {
                "row_id": gold["item_id"],
                "participant": run["participant"],
                "condition": run["condition"],
                "round": run["round"],
                "case": gold["case"],
                "level": "Evidence item",
                "target_id": gold["item_id"],
                "topic_code": gold["topic_code"],
                "gold_decision": gold["gold_action"],
                "actual_decision": actual["actual_action"],
                "correct": correct,
                "issue_bearing": issue_bearing,
                "issue_recognized": correct if issue_bearing else "",
                "missingness_expected": gold["gold_action"] == "edit_and_flag_partial",
                "missingness_actual": actual.get("actual_missingness", False),
                "missingness_correct": "",
                "false_approval": is_false_approval(actual, gold),
                "source_file": actual["source_file"],
            }
        )

    objective_rows.sort(
        key=lambda row: (
            row["participant"],
            row["round"],
            0 if row["level"] == "Topic" else 1,
            row["target_id"],
        )
    )
    if len(objective_rows) != 60 or len({row["row_id"] for row in objective_rows}) != 60:
        raise ValueError("The rebuilt objective table is not a unique 60-row universe")

    condition_scores: list[dict[str, Any]] = []
    for run in responses:
        rows = [row for row in objective_rows if row["case"] == run["case"]]
        topic_rows = [row for row in rows if row["level"] == "Topic"]
        item_rows = [row for row in rows if row["level"] == "Evidence item"]
        issue_rows = [row for row in item_rows if row["issue_bearing"]]
        expected_missing_rows = [row for row in topic_rows if row["missingness_expected"]]
        condition_scores.append(
            {
                "participant": run["participant"],
                "name": run["name"],
                "round": run["round"],
                "condition": run["condition"],
                "case": run["case"],
                "protocol": run["protocol"],
                "topic_correct": sum(bool(row["correct"]) for row in topic_rows),
                "topic_total": len(topic_rows),
                "item_correct": sum(bool(row["correct"]) for row in item_rows),
                "item_total": len(item_rows),
                "issues_recognized": sum(bool(row["issue_recognized"]) for row in issue_rows),
                "issues_total": len(issue_rows),
                "missingness_correct": sum(bool(row["missingness_correct"]) for row in expected_missing_rows),
                "missingness_total": len(expected_missing_rows),
                "false_approvals": sum(bool(row["false_approval"]) for row in item_rows),
                "PC1": run.get("PC1"),
                "PC2": run.get("PC2"),
                "PC3": run.get("PC3"),
                "PC4": run.get("PC4"),
                "PC5": run.get("PC5"),
                "PC6": run.get("PC6"),
                "PC7": run.get("PC7"),
                "PC8": run.get("PC8"),
                "PC9": run.get("PC9"),
                "PC10": run.get("PC10"),
                "PC11": run.get("PC11"),
                "source_file": run["source_member"],
            }
        )
    condition_scores.sort(key=lambda row: (row["participant"], row["round"]))

    aggregate: dict[str, dict[str, Any]] = {}
    for condition in ("Manual", "Prototype"):
        rows = [row for row in condition_scores if row["condition"] == condition]
        aggregate[condition] = {
            "runs": len(rows),
            "topic_correct": sum(row["topic_correct"] for row in rows),
            "topic_total": sum(row["topic_total"] for row in rows),
            "item_correct": sum(row["item_correct"] for row in rows),
            "item_total": sum(row["item_total"] for row in rows),
            "issues_recognized": sum(row["issues_recognized"] for row in rows),
            "issues_total": sum(row["issues_total"] for row in rows),
            "missingness_correct": sum(row["missingness_correct"] for row in rows),
            "missingness_total": sum(row["missingness_total"] for row in rows),
            "false_approvals": sum(row["false_approvals"] for row in rows),
            "time_values_minutes": [float(row["PC11"]) for row in rows],
            "time_mean_minutes": round(mean(float(row["PC11"]) for row in rows), 1),
            "time_median_minutes": statistics.median(float(row["PC11"]) for row in rows),
            "mental_effort_values": [float(row["PC8"]) for row in rows],
            "mental_effort_median": statistics.median(float(row["PC8"]) for row in rows),
        }

    paired_time_differences: list[float] = []
    for participant in FORMAL_PARTICIPANTS:
        manual = next(row for row in condition_scores if row["participant"] == participant and row["condition"] == "Manual")
        prototype = next(row for row in condition_scores if row["participant"] == participant and row["condition"] == "Prototype")
        paired_time_differences.append(float(manual["PC11"]) - float(prototype["PC11"]))

    comparisons: list[dict[str, Any]] = []
    for participant in FORMAL_PARTICIPANTS:
        comparison_source = next(row for row in responses if row["participant"] == participant and row["comparison"])
        comparisons.append(
            {
                "participant": participant,
                "name": PARTICIPANT_NAMES[participant],
                **comparison_source["comparison"],
                "source_file": comparison_source["source_member"],
            }
        )

    analysis = {
        "study_id": "PS-TRV-2026-01",
        "source_date": "2026-08-11",
        "formal_sample_n": len(FORMAL_PARTICIPANTS),
        "formal_participants": list(FORMAL_PARTICIPANTS),
        "pilot_in_formal_statistics": False,
        "time_basis": "PC11 objectively timed active review minutes; descriptive n=3 comparison",
        "condition_runs": responses,
        "condition_scores": condition_scores,
        "objective_rows": objective_rows,
        "aggregate": aggregate,
        "paired_time_difference_values_minutes": paired_time_differences,
        "paired_time_difference_mean_minutes": round(mean(paired_time_differences), 1),
        "comparisons": comparisons,
        "prototype_preference_count": sum(clean_text(row.get("D1")) == "Prototype" for row in comparisons),
        "prototype_preference_total": len(comparisons),
        "input_hashes": input_hashes,
        "interpretation_boundary": (
            "Exploratory n=3 results from controlled synthetic protocols; no claim of statistical significance, "
            "general effectiveness, production readiness, or generalisable large-scale efficiency."
        ),
    }
    return analysis


OBJECTIVE_COLUMNS = [
    "row_id",
    "participant",
    "condition",
    "round",
    "case",
    "level",
    "target_id",
    "topic_code",
    "gold_decision",
    "actual_decision",
    "correct",
    "issue_bearing",
    "issue_recognized",
    "missingness_expected",
    "missingness_actual",
    "missingness_correct",
    "false_approval",
    "source_file",
]


def summary_text(analysis: dict[str, Any]) -> str:
    manual = analysis["aggregate"]["Manual"]
    prototype = analysis["aggregate"]["Prototype"]
    return "\n".join(
        [
            "重建结果摘要 / Rebuilt Result Summary",
            "",
            f"Formal n={analysis['formal_sample_n']}",
            f"- Objective rows: {len(analysis['objective_rows'])} (24 Topic + 36 Evidence-item).",
            f"- Topic correct: Manual {manual['topic_correct']}/{manual['topic_total']}; Prototype {prototype['topic_correct']}/{prototype['topic_total']}.",
            f"- Evidence-item correct: Manual {manual['item_correct']}/{manual['item_total']}; Prototype {prototype['item_correct']}/{prototype['item_total']}.",
            f"- E3-E6 issues detected: Manual {manual['issues_recognized']}/{manual['issues_total']}; Prototype {prototype['issues_recognized']}/{prototype['issues_total']}.",
            f"- Missingness correct: Manual {manual['missingness_correct']}/{manual['missingness_total']}; Prototype {prototype['missingness_correct']}/{prototype['missingness_total']}.",
            f"- False approvals: Manual {manual['false_approvals']}; Prototype {prototype['false_approvals']}.",
            f"- PC11 objectively timed active review time: Manual mean {manual['time_mean_minutes']:.1f} min, median {manual['time_median_minutes']:.0f}; Prototype mean {prototype['time_mean_minutes']:.1f} min, median {prototype['time_median_minutes']:.0f}; mean paired difference {analysis['paired_time_difference_mean_minutes']:.1f} min.",
            f"- PC8 mental effort median: Manual {manual['mental_effort_median']:.0f}; Prototype {prototype['mental_effort_median']:.0f}.",
            f"- Condition preference: Prototype {analysis['prototype_preference_count']}/{analysis['prototype_preference_total']}.",
            "",
            "Interpretation boundary",
            analysis["interpretation_boundary"],
            "",
        ]
    )


def write_outputs(analysis: dict[str, Any], output_dir: Path, package_root: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    objective_path = output_dir / "Objective_60_Rows_REBUILT.csv"
    with objective_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OBJECTIVE_COLUMNS)
        writer.writeheader()
        writer.writerows(analysis["objective_rows"])

    with (output_dir / "Analysis_Data_REBUILT.json").open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(analysis, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")

    (output_dir / "Result_Summary_REBUILT.txt").write_text(summary_text(analysis), encoding="utf-8", newline="\n")

    with (output_dir / "Input_SHA256.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["input_type", "participant", "path", "sha256"])
        writer.writeheader()
        writer.writerows(analysis["input_hashes"])

    run_log_lines = [
        "Analysis rebuild run log",
        f"run_timestamp_utc: {datetime.now(timezone.utc).isoformat()}",
        f"python_version: {sys.version.split()[0]}",
        f"openpyxl_version: {openpyxl.__version__}",
        "package_root: .",
        f"formal_sample_n: {analysis['formal_sample_n']}",
        f"condition_runs: {len(analysis['condition_runs'])}",
        f"objective_rows: {len(analysis['objective_rows'])}",
        f"pilot_in_formal_statistics: {analysis['pilot_in_formal_statistics']}",
        "status: PASS",
        "",
        "Input SHA-256",
    ]
    run_log_lines.extend(f"{row['sha256']}  {row['path']}" for row in analysis["input_hashes"])
    (output_dir / "Analysis_Run_Log.txt").write_text("\n".join(run_log_lines) + "\n", encoding="utf-8", newline="\n")


def main() -> int:
    default_package_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package-root", type=Path, default=default_package_root)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "outputs")
    args = parser.parse_args()
    analysis = build_analysis(args.package_root)
    write_outputs(analysis, args.output_dir, args.package_root)
    print(summary_text(analysis), end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
